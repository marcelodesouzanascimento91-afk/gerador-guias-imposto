"""Relatórios em Excel (processamento e consultas por período)."""
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = PatternFill("solid", start_color="1F3864")
CINZA = PatternFill("solid", start_color="F2F2F2")
AMARELO = PatternFill("solid", start_color="FFF2CC")
VERMELHO = PatternFill("solid", start_color="F8CBAD")
FONTE = "Arial"
BORDA = Border(bottom=Side(style="thin", color="BFBFBF"))
MOEDA = 'R$ #,##0.00;[Red]-R$ #,##0.00;"-"'

COLUNAS = [
    ("Nº NF", "nf", 10, None), ("Cliente", "cliente", 38, None), ("Tipo de Imposto", "tipo", 14, None),
    ("Valor", "valor", 14, MOEDA), ("Vencimento", "vencimento", 12, "DD/MM/YYYY"),
    ("UF", "uf", 6, None), ("Receita", "receita", 9, None), ("Multa (est.)", "multa", 12, MOEDA),
    ("Juros (est.)", "juros", 12, MOEDA), ("Total", None, 14, MOEDA),
    ("Data Pagamento", "data_pagamento", 13, "DD/MM/YYYY"), ("Dias Atraso", "dias_atraso", 9, "0"),
    ("Situação", "situacao", 32, None), ("Modalidade", "modo", 30, None), ("Série", "serie", 6, None),
    ("Emissão NF", "data_emissao", 12, "DD/MM/YYYY"), ("CNPJ/CPF Cliente", "cliente_doc", 17, "@"),
    ("Venc. Original", "vencimento_original", 12, "DD/MM/YYYY"),
    ("Linha Digitável", "linha_digitavel", 52, "@"), ("Nº Controle", "nosso_numero", 16, "@"),
    ("Arquivo PDF", "arquivo_pdf", 45, None), ("Chave de Acesso", "chave", 47, "@"),
    ("Observação", "observacao", 70, None),
]


def _data(v):
    if not v:
        return None
    if isinstance(v, (date, datetime)):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return v


def _cabecalho(ws, linha, titulos, larguras):
    for i, (t, w) in enumerate(zip(titulos, larguras), 1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
        c.fill = AZUL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[linha].height = 28


def _informativo(row):
    return (row.get("situacao") or "") == "INFORMATIVO" or (row.get("modo") or "") == "Informativo"


def exportar(rows, caminho, titulo="Relatório de Guias de Impostos", ocorrencias=None, filtros=""):
    informativos = [x for x in rows if _informativo(x)]
    rows = [x for x in rows if not _informativo(x)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Guias"
    ws["A1"] = titulo
    ws["A1"].font = Font(name=FONTE, bold=True, size=14, color="1F3864")
    ws["A2"] = f"Gerado em {datetime.now():%d/%m/%Y %H:%M}" + (f"  |  Filtros: {filtros}" if filtros else "")
    ws["A2"].font = Font(name=FONTE, italic=True, size=9, color="595959")
    ws["A3"] = "Multa e juros são estimativas; o valor oficial é o informado na guia emitida pelo portal/SEFAZ."
    ws["A3"].font = Font(name=FONTE, italic=True, size=8, color="808080")
    hl = 5
    _cabecalho(ws, hl, [c[0] for c in COLUNAS], [c[2] for c in COLUNAS])
    col = {c[0]: get_column_letter(i) for i, c in enumerate(COLUNAS, 1)}
    r = hl
    for row in rows:
        r += 1
        for i, (tit, chave, _, fmt) in enumerate(COLUNAS, 1):
            if tit == "Total":
                val = f"={col['Valor']}{r}+{col['Multa (est.)']}{r}+{col['Juros (est.)']}{r}"
            else:
                val = row.get(chave, "")
                if fmt == "DD/MM/YYYY":
                    val = _data(val)
                elif chave == "nf" and str(val).isdigit():
                    val = int(val)
            c = ws.cell(row=r, column=i, value=val)
            c.font = Font(name=FONTE, size=9)
            c.border = BORDA
            if fmt:
                c.number_format = fmt
        sit = (row.get("situacao") or "").upper()
        if sit in ("VENCIDA",) or "REJEIT" in sit or "ERRO" in sit:
            for i in range(1, len(COLUNAS) + 1):
                ws.cell(row=r, column=i).fill = VERMELHO if ("REJEIT" in sit or "ERRO" in sit) else AMARELO
    ultima = r
    # Totais
    r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = Font(name=FONTE, bold=True)
    for tit in ("Valor", "Multa (est.)", "Juros (est.)", "Total"):
        c = ws[f"{col[tit]}{r}"]
        c.value = f"=SUBTOTAL(9,{col[tit]}{hl + 1}:{col[tit]}{ultima})" if ultima > hl else 0
        c.font = Font(name=FONTE, bold=True)
        c.number_format = MOEDA
        c.fill = CINZA
    ws.freeze_panes = ws.cell(row=hl + 1, column=3)
    if ultima > hl:
        ws.auto_filter.ref = f"A{hl}:{get_column_letter(len(COLUNAS))}{ultima}"

    # Resumo por tipo e UF (fórmulas ligadas à aba Guias)
    rs = wb.create_sheet("Resumo")
    rs["A1"] = "Resumo por tipo de imposto e UF"
    rs["A1"].font = Font(name=FONTE, bold=True, size=13, color="1F3864")
    _cabecalho(rs, 3, ["Tipo de Imposto", "UF", "Qtde", "Valor", "Multa + Juros (est.)", "Total"], [18, 6, 8, 16, 18, 16])
    pares = sorted({(x.get("tipo", ""), x.get("uf", "")) for x in rows})
    rr = 3
    faixa = lambda c: f"Guias!${col[c]}${hl + 1}:${col[c]}${max(ultima, hl + 1)}"
    for tipo, uf in pares:
        rr += 1
        rs.cell(row=rr, column=1, value=tipo)
        rs.cell(row=rr, column=2, value=uf)
        cond = f'{faixa("Tipo de Imposto")},A{rr},{faixa("UF")},B{rr}'
        rs.cell(row=rr, column=3, value=f"=COUNTIFS({cond})")
        rs.cell(row=rr, column=4, value=f"=SUMIFS({faixa('Valor')},{cond})").number_format = MOEDA
        rs.cell(row=rr, column=5, value=f"=SUMIFS({faixa('Multa (est.)')},{cond})+SUMIFS({faixa('Juros (est.)')},{cond})").number_format = MOEDA
        rs.cell(row=rr, column=6, value=f"=D{rr}+E{rr}").number_format = MOEDA
        for c in range(1, 7):
            rs.cell(row=rr, column=c).font = Font(name=FONTE, size=10)
    rr += 1
    rs.cell(row=rr, column=1, value="TOTAL").font = Font(name=FONTE, bold=True)
    for c, L in ((3, "C"), (4, "D"), (5, "E"), (6, "F")):
        cell = rs.cell(row=rr, column=c, value=f"=SUM({L}4:{L}{rr - 1})" if rr > 4 else 0)
        cell.font = Font(name=FONTE, bold=True)
        cell.fill = CINZA
        if c > 3:
            cell.number_format = MOEDA

    if informativos:
        inf = wb.create_sheet("Informativos")
        inf["A1"] = "Situações que não geram GNRE (ST interna, SUFRAMA) – conferência"
        inf["A1"].font = Font(name=FONTE, bold=True, size=13, color="1F3864")
        cols_inf = [("Nº NF", "nf", 10), ("Cliente", "cliente", 38), ("Tipo", "tipo", 14), ("UF", "uf", 6),
                    ("Valor de referência", "valor", 16), ("Emissão NF", "data_emissao", 12),
                    ("Orientação", "observacao", 110), ("Chave de Acesso", "chave", 47)]
        _cabecalho(inf, 3, [c[0] for c in cols_inf], [c[2] for c in cols_inf])
        for i, row in enumerate(informativos, 4):
            for j, (_, k, _) in enumerate(cols_inf, 1):
                v = row.get(k, "")
                if k == "data_emissao":
                    v = _data(v)
                elif k == "nf" and str(v).isdigit():
                    v = int(v)
                c = inf.cell(row=i, column=j, value=v)
                c.font = Font(name=FONTE, size=9)
                if k == "valor":
                    c.number_format = MOEDA
                elif k == "data_emissao":
                    c.number_format = "DD/MM/YYYY"
                elif k == "chave":
                    c.number_format = "@"
                elif k == "observacao":
                    c.alignment = Alignment(wrap_text=True, vertical="top")

    if ocorrencias is not None:
        oc = wb.create_sheet("Ocorrências")
        oc["A1"] = "Erros, avisos e verificações encontrados nos XMLs"
        oc["A1"].font = Font(name=FONTE, bold=True, size=13, color="1F3864")
        _cabecalho(oc, 3, ["Nível", "Nº NF", "Mensagem", "Chave de Acesso", "Arquivo"], [9, 10, 100, 47, 50])
        for i, o in enumerate(ocorrencias, 4):
            vals = [o.nivel, int(o.nNF) if str(o.nNF).isdigit() else o.nNF, o.mensagem, o.chave, o.arquivo]
            for j, v in enumerate(vals, 1):
                c = oc.cell(row=i, column=j, value=v)
                c.font = Font(name=FONTE, size=9, bold=(j == 1))
                c.alignment = Alignment(wrap_text=(j == 3), vertical="top")
            if o.nivel == "ERRO":
                oc.cell(row=i, column=1).fill = VERMELHO
            elif o.nivel == "AVISO":
                oc.cell(row=i, column=1).fill = AMARELO
        oc.freeze_panes = "A4"
        if ocorrencias:
            oc.auto_filter.ref = f"A3:E{len(ocorrencias) + 3}"
    wb.save(caminho)
    return caminho
